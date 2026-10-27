"""Build a NEW Korean thesis after actual372 numerical QA and explicit ROOT GO.

No original document edits, raw inputs, NN, PnP, bootstrap, or model selection.
The original18page outline is source prose, not a required new page count.
Use bundled Python. Rendering uses the old task-local compatibility wrapper
around the unchanged bundled LibreOffice binary, never desktop/system LO.
"""
from pathlib import Path
from collections import Counter
import argparse
import copy
import json
import re
import subprocess

import publication_common as P
import prepare_publication as F

RUNTIME=Path('/home/minjae/.cache/codex-runtimes/codex-primary-runtime/dependencies')
PYTHON=RUNTIME/'python/bin/python3';NODE=RUNTIME/'node/bin/node'
MARKER=Path('/home/minjae/.codex/plugins/cache/openai-primary-runtime/documents/26.1017.31632/skills/documents/container_tools/mark_artifact_operation_started.mjs')
PDF_MARKER=Path('/home/minjae/.codex/plugins/cache/openai-primary-runtime/pdf/26.1017.31632/skills/pdf/container_tools/mark_artifact_operation_started.mjs')
BUNDLED_LO=RUNTIME/'native/libreoffice-headless/libreoffice/program/soffice.bin'
TITLE='실제 파렛트 영상의 의사 레이블을 이용한 DOPE 자기학습 성능 평가'
LABEL={'R0':'기본','control':'합성 대조','full':'전체 필터','no_flip':'반전 제외','no_confidence':'신뢰도 제외',
    'no_ransac':'RANSAC 제외','no_loo':'LOO 제외','no_final_support':'최종 지지점 문턱 제외','confidence_only':'신뢰도만'}
MAIN_CONDITIONS={'R0','control','full','no_flip'}


def decision_ko(value):
    labels={'candidate_higher':'높았다','full_higher':'높았다','candidate_lower':'낮았다','full_lower':'낮았다',
        'includes_zero':'차이 불명확','NA_model_unavailable':'계산 불가',
        'NA_undefined_denominator_no_draws_dropped':'계산 불가'}
    P.require(value in labels,'Unknown prespecified CI decision');return labels[value]


def exposure_summary(io):
    rows=io.csv(io.one('ratio_exposure_summary'))
    wanted={(k,c,b) for k in P.EXPECTED for c in ('A','B','C','D','E') for b in ('real','synthetic')}
    keys=[(r['kind'],r['condition'],r['branch']) for r in rows]
    P.require(len(rows)==len(set(keys))==30 and set(keys)==wanted,'Exact30 ratio exposure branch rows required')
    for row in rows:
        P.require(int(row['total_draws'])==6624 and int(row['draws_per_update'])==8
            and int(row['n_updates'])==828 and int(row['epochs'])==3,'Committed R2 exposure budget differs')
    return rows


def prose_copy(block):
    row=copy.deepcopy(block)
    if row['type']=='p':
        text=row['text']
        for word,label in (('teacher','교사 모델'),('student','학생 모델'),('flip','수평 반전')):
            text=re.sub(r'(?<![A-Za-z0-9_])'+word+r'(?![A-Za-z0-9_])',label,text)
        for label in ('교사 모델','학생 모델'):
            for before,after in (('를','을'),('는','은'),('가','이'),('와','과')):text=text.replace(label+before,label+after)
        row['text']=text
    return row


def test_prose_copy_fixtures():
    """Pure fixtures prepared for ROOT; never called during build/import."""
    examples=[
        ('반전 좌표는 x = W − 1 − xflip으로 복원한다.','반전 좌표는 x = W − 1 − xflip으로 복원한다.'),
        ('Noisy Student는 teacher의 의사 레이블을 활용하여 student를 학습한다.',
         'Noisy Student는 교사 모델의 의사 레이블을 활용하여 학생 모델을 학습한다.'),
        ('teacher는 student가 teacher를 학습시키는 예시이다.','교사 모델은 학생 모델이 교사 모델을 학습시키는 예시이다.'),
        ('LOO와 flip 필터를 사용하며 flip은 두 영상 추론 조건이다.','LOO와 수평 반전 필터를 사용하며 수평 반전은 두 영상 추론 조건이다.'),
        ('LOO 및 flip의 0.05는 D로 나눈 값이다.','LOO 및 수평 반전의 0.05는 D로 나눈 값이다.'),
        ('xflip flip_R1 R1_flip teacher_model student2 preteacher teacher_student',
         'xflip flip_R1 R1_flip teacher_model student2 preteacher teacher_student')]
    for original,expected in examples:
        block={'type':'p','text':original};result=prose_copy(block)
        assert result['text']==expected and block['text']==original and result is not block
    equation={'type':'equation','kind':'loss'};assert prose_copy(equation)==equation
    return {'passed':True,'cases':len(examples)+1,'data_reads':0,'authoring':0}


def marker(fmt):
    subprocess.run([str(NODE),str(MARKER if fmt=='docx' else PDF_MARKER),'--operation-kind','create','--expected-output-count','1','--output-format',fmt],check=True)


def check_sources(go):
    required={str(Path(__file__).resolve()),str(Path(P.__file__).resolve()),str(Path(F.__file__).resolve()),str(P.OUT/'code/final_numerical_qa.py')}
    P.require(required.issubset(go['source_sha256']),'Builder and all numerical/content helpers must be frozen')
    for p,digest in go['source_sha256'].items():
        P.require(Path(p).is_relative_to(P.OUT/'code') and P.sha(p)==digest,'Manuscript code source changed')
    P.require(go['source_sha256'][str(Path(__file__).resolve())]==go['builder_sha256'],'Builder source authorities disagree')


def manuscript_authority(io,receipt=None):
    path=P.OUT/'preflight/MANUSCRIPT_GO.json';go=json.loads(path.read_text());digest=P.sha(path)
    P.require(go.get('authorized') is True and go.get('action')=='build_updated_Korean_DOCX_PDF'
        and go['manuscript_inputs_sha256']==io.manifest_sha and go['numerical_inputs_sha256']==P.sha(P.NUMERICAL_INPUTS)
        and go['builder_sha256']==P.sha(Path(__file__).resolve()),'ROOT manuscript GO/input authority differs')
    check_sources(go)
    if receipt is not None:
        P.require(receipt['ROOT_GO_sha256']==digest and receipt['builder_sha256']==go['builder_sha256']
            and receipt['code_source_sha256']==go['source_sha256']
            and receipt['source_sha256'][str(P.MANUSCRIPT_INPUTS)]==io.manifest_sha,'Built manuscript authority changed')
        for source,bound in receipt['source_sha256'].items():
            P.require(P.sha(source)==bound and (source==str(P.MANUSCRIPT_INPUTS) or source in io.files
                and io.files[source]['sha256']==bound),'Built manuscript source/input contract changed')
    return path,go,digest


def content(io):
    qa_path,_=F.numerical_ready(io);points,pose,ratio,ci,ratio_ci,teacher=F.tables(io)
    old=io.read(io.one('manuscript_seed'));P.require(len(old['pages'])==18 and len(old['references'])==6,'Verified original draft/reference seed differs')
    for path in io.role('methods_origin_evidence'):io.path(path)
    metadata=io.read(io.one('evaluation_manifest_metadata'))['records'];P.require(len(metadata)==372,'Exact372 metadata required')
    resolution={k:' / '.join(f'{w}×{h} {n}장' for (w,h),n in sorted(Counter((f['width'],f['height']) for f in metadata if f['kind']==k).items())) for k in P.EXPECTED}
    main_ci=io.csv(io.one('ci_main_comparisons'));exposures=exposure_summary(io)
    sections=[]
    def chapter(title):
        row={'title':title,'blocks':[]};sections.append(row);return row['blocks']
    def paragraph(blocks,text):blocks.append({'type':'p','text':text})
    def table(blocks,caption,headers,rows,widths=None):blocks.append({'type':'table','caption':caption,'headers':headers,'rows':rows,'widths':widths})
    lookup={(r['kind'],r['condition'],int(r['round'])):r for r in points}
    def result_phrase(kind):
        a,b=lookup[kind,'R0',0],lookup[kind,'full',2]
        if b['available']!='True':return f'{kind}의 전체 필터 R2는 학습을 완료한 모델이 없어 계산 불가로 보고하였다.'
        return f"{kind}의 PCK10은 기본 모델 R0에서 {F.fmt(a['PCK10'],100)}%, 전체 필터 R2에서 {F.fmt(b['PCK10'],100)}%였다."
    abstract=chapter('국문초록')
    paragraph(abstract,'합성 영상으로 학습한 DOPE의 실제 영상 적응 성능을 확인하기 위해 세 종류의 파렛트에 자기학습을 각각 적용하였다. 신뢰도, RANSAC, 한 점을 제외한 자세 재추정(LOO), 수평 반전 일관성으로 기준점을 선별하고, 승인된 원예측 좌표를 학습 목표로 사용하였다. 업데이트 수를 맞춘 합성 추가학습 대조군과 필터 구성 요소 제거 조건을 비교하였다.')
    paragraph(abstract,'주평가는 PLASTIC 128장, WOOD 125장, GREEN 119장으로 구성된 총 372장과 2,891개 참조점에서 수행하였다. '+' '.join(result_phrase(k) for k in P.EXPECTED))
    paragraph(abstract,'실제·합성 고유 자료 수의 비율을 달리한 15개 R2 조건을 추가로 비교하였으며, 손실계수는 고정하였다. 프레임 단위 신뢰구간을 주분석으로 사용하고, 소수 녹화 단위 신뢰구간은 보조 분석으로 구분하였다. 2D 기준점으로 재구성한 참조에 대한 자세 오차는 독립적으로 실측한 6D 정확도를 의미하지 않는다. 본 결과는 단일 학습 시드와 반복 사용한 평가 자료에 기반하며, 종류 및 조건별 차이를 구분하여 해석하였다.')
    paragraph(abstract,'주제어: DOPE, 파렛트, 자기학습, 의사 레이블, 필터 제거, 자료 수 비율')
    intro=chapter('제1장 서론')
    for b in old['pages'][3]['blocks']:
        if b['type']=='h2' and b['text'].startswith('3 '):break
        if b['type']=='p' and '연구 범위는 보유 자료' in b['text']:
            paragraph(intro,'학습에는 평가 주석을 사용하지 않았다. 평가 자료를 반복 사용하였고 참조점의 생성 출처가 혼합되어 있으므로, 2D 기준점 비교를 주평가로 정하였다. 재구성 참조에 대한 6D 지표는 보조 분석으로 구분하였다.')
        else:intro.append(prose_copy(b))
    paragraph(intro,'제2장에서는 관련 연구를 정리하고, 제3장에서는 자기학습 방법과 비교 조건을 설명한다. 제4장에서는 실험 자료와 결과를 제시한다. 제5장에서는 결과의 해석과 한계를 논의하고, 제6장에서 결론을 제시한다.')
    related=chapter('제2장 관련 연구');related.extend(prose_copy(b) for b in old['pages'][4]['blocks'])
    methods=chapter('제3장 자기학습 방법')
    for page in old['pages'][5:8]:
        for b in page['blocks']:
            if b['type']=='figure' or (b['type']=='equation' and b['kind']=='loss'):continue
            if b['type']=='p' and b['text'].startswith('주 실험은 LOO와 flip'):
                paragraph(methods,'전체 필터에서 신뢰도, RANSAC, LOO, 반전 일관성, 최종 지지점 수 문턱을 각각 제거한 조건과 신뢰도만 사용하는 조건을 비교하였다. 종류별로 사전 고정한 R1과 R2를 모두 보고하였으며, 평가 결과에 따라 차수나 문턱을 선택하지 않았다. 필터 통과 여부가 달라지므로 조건별 의사 레이블의 지지 집합은 서로 다를 수 있다.');continue
            methods.append(prose_copy(b))
    paragraph(methods,'전체 필터는 신뢰도 0.3, 원본 영상에서 5 px인 RANSAC 문턱, LOO 및 반전 오차 0.05D, 최소 6개의 지지점을 사용하였다. 각 제거 조건은 사전 고정한 코드 정의를 따랐다. RANSAC 제외 조건에서는 사용 가능한 점 전체에 PnP를 적합하였다. 최종 지지점 문턱 제외 조건에서는 프레임의 최종 승인 문턱만 완화하고, 기하 적합에 필요한 최소 6점 조건은 유지하였다. 신뢰도 제외 조건은 학습용 예측의 상류 복호화 문턱까지 제거한 조건이다. 평가에서는 모든 모델에 동일한 신뢰도 0.3 복호화를 적용하였다.')
    methods.append({'type':'equation','kind':'loss'})
    paragraph(methods,'전체 손실은 합성 분기 평균 손실과 실제 분기 평균 손실의 합으로 구성하고 λ=1로 고정하였다. 자료 수 비율 실험에서는 고유 자료의 수를 변경하였으며, 자료 수에 따라 손실 가중치나 λ를 변경하지 않았다. 필터 제거 주실험은 실제 풀의 크기가 작을 때 짧은 배치를 사용하는 원래 정책을 따랐다. 반면 자료 수 비율 실험은 독립적인 순환 추출로 실제·합성 분기를 각각 8회 사용분으로 채우는 cycle-fill 정책을 사용하였다. 반복 노출과 고유 영상 수를 구분하였으며, 이러한 배치 정책의 차이를 고려하여 모든 실험의 실제 영상 노출 수가 같다고 해석하지 않았다. R2 학습은 R1의 모델, Adam 및 scaler 상태를 이어서 수행하였다.')
    paragraph(methods,'R0의 원학습에서는 640×480 영상을 400×400으로 비등방 축소하였다. 현재 추가학습 및 평가에서는 100 px 반사 패딩 후 원래 크기로 조정하고, 종횡비를 유지하면서 짧은 변을 400에 맞추었다. 이때 640×480 영상의 입력 크기는 약 400×528이 된다. 현재의 11×11 가중평균 복호화와 출력 격자 0.4395칸 오프셋은 과거의 7×7 softargmax 및 T=0.1 경로와 다르다. R0에도 현재 평가 파이프라인을 적용하여 출력 비교 규약을 통일하였으나, 이는 원학습의 전처리와 복호화를 그대로 재현하였다는 의미는 아니다. 이 차이는 기존 방법 및 복호화 감사 기록에 근거하여 기술하였다.')
    results=chapter('제4장 실험 자료와 결과')
    results.append({'type':'h2','text':'4.1 평가 자료와 지표'})
    table(results,'통합 평가 모집단',['종류','영상 수','참조점','수동 보조점','원본 해상도'],
        [['PLASTIC',128,985,'없음',resolution['PLASTIC']],['WOOD',125,963,330,resolution['WOOD']],['GREEN',119,943,600,resolution['GREEN']]], [.9,.65,.65,.8,3.2])
    paragraph(results,'목재 80장은 원자료의 FINAL 역할을 보존한 상태로 기존 45장에 추가하여 통합 평가하였다. 추가 영상은 학습 풀과 녹화 REC001/002를 공유하므로, 녹화가 독립된 시험 자료로 볼 수 없다. 또한 동일한 원본 픽셀 문턱을 사용하더라도 해상도가 다르면 정규화된 허용오차가 달라진다. 기존 45장에서 얻은 PCK10 60.12%와 새 125장의 결과 간 차이는 평가 모집단의 변화를 포함하므로 모델의 개선량으로 해석하지 않았다. 모델 간 차이는 동일한 125장에 짝지어 비교하였다.')
    results.append({'type':'equation','kind':'pck'})
    paragraph(results,'Ncorrect는 오차가 문턱 τ 이내인 점의 수이고, Nreference는 전체 참조점 수이다. 미검출은 실패로 처리하고 중심점인 채널 8은 채점에서 제외하였다. 관측률(Coverage)은 전체 참조점 중 예측이 관측된 점의 비율이며, 조건부 정밀도(conditional_precision10)는 관측된 참조점 중 오차가 10 px 이내인 점의 비율이다. 오차의 평균, 분산, 표준편차(SD), 중앙값은 관측된 점에 조건부로 계산하였다.')
    results.append({'type':'h2','text':'4.2 기본 모델과 자기학습 비교'})
    for kind in P.EXPECTED:
        rows=[r for r in points if r['kind']==kind and r['condition'] in MAIN_CONDITIONS]
        P.require(len(rows)==7,'Expected seven main roles per kind')
        table(results,kind+' 기본 비교 모델의 2D 지표',['모델','R','PCK5 %','PCK10 %','PCK20 %','관측률 %','오차 평균±SD px','오차 n'],
            [[LABEL[r['condition']],r['round'],*[F.fmt(r[k],100) for k in ('PCK5','PCK10','PCK20','coverage')],F.fmt(r['conditional_pixel_error_mean'])+' ± '+F.fmt(r['conditional_pixel_error_std']),r['conditional_pixel_error_n']] for r in rows],
            [1.1,.3,.6,.65,.65,.65,1.6,.65])
    actual_checkpoints=len({r['checkpoint_sha256'] for r in points if r['checkpoint_sha256']})
    paragraph(results,f'종류별 17개씩 총 51개 평가조합을 분석하였다. 여러 종류가 공유한 R0와 합성 대조 R1/R2를 중복으로 세지 않으면, 학습을 완료한 고유 체크포인트는 {actual_checkpoints}개이다. 전체 지표와 분산·사분위수, WOOD 수동 330점 및 GREEN 수동 600점 패널의 결과는 CSV에 보존하였다. 본문에서는 기본 모델, 합성 대조군, 전체 필터, 반전 제외 조건에 해당하는 종류별 7개씩 총 21개 평가조합을 제시한다. 사전 고정한 모든 차수를 보고하였으며, 평가 결과에 따라 차수나 체크포인트를 선택하지 않았다.')
    for kind in P.EXPECTED:
        rr=[r for r in main_ci if r['kind']==kind and r['metric']=='PCK10']
        table(results,kind+' 기본 비교 모델 간 PCK10 차이',['비교 모델','기준 모델','차이 %p','95% CI %p','판정'],
            [[LABEL[r['candidate_condition']]+' R'+r['candidate_round'],LABEL[r['comparator_condition']]+' R'+r['comparator_round'],F.fmt(r['point_difference'],100),F.fmt(r['ci95_low'],100)+' ~ '+F.fmt(r['ci95_high'],100),decision_ko(r['decision'])] for r in rr],[1.25,1.25,.8,1.6,1.3])
    results.append({'type':'h2','text':'4.3 필터 구성 요소 제거 실험'})
    for kind in P.EXPECTED:
        rr=[r for r in ci if r['kind']==kind and r['metric']=='PCK10']
        table(results,kind+' 전체 필터에서 제거 조건을 뺀 PCK10 차이',['차수','제거 조건','차이 %p','95% CI %p','판정'],
            [[r['round'],LABEL[r['comparator']],F.fmt(r['point_difference_full_minus_ablation'],100),F.fmt(r['ci95_low'],100)+' ~ '+F.fmt(r['ci95_high'],100),decision_ko(r['decision'])] for r in rr],[.45,1.8,.8,1.6,1.55])
    paragraph(results,'짝지은 프레임 부트스트랩은 시드 20261026으로 10,000회 수행하였으며, 점 수 합계의 비율 차이를 통계량으로 사용하였다. 종류별로 모델, 차수, 지표 및 수동 패널에 동일한 추출 가중치를 공유하였다. 표의 판정은 95% 신뢰구간(CI)의 하한이 0보다 크면 높았다, 상한이 0보다 작으면 낮았다, 구간이 0을 포함하면 차이 불명확으로 표시하였다. 원 CSV의 판정 문자열은 그대로 보존하였다. 차이 불명확은 동등성이나 효과가 없다는 증거가 아니다. 분모가 0인 추출이 하나라도 있으면 해당 CI 전체를 계산 불가로 두고 재추출하지 않았다. 목재의 녹화 단위 CI는 4개 군집에 기반한 보조 분석이며, 프레임 단위 CI를 주결과로 유지하였다.')
    paragraph(results,'동일한 R0 교사 예측에 각 필터를 적용한 선별 정밀도, 올바른 점 유지율, 잘못된 점 승인율은 COMPONENT_QUALITY.csv에 제시하였다. 여기에는 신뢰도만 사용하는 조건도 포함된다. 단계별 선별·탈락 기준점의 통계는 T3_FILTER_STAGES.csv와 5/10/20 px 상세표에 제시하였다. 이 진단은 필터 자체의 선별 효과를 나타낸다. 선별된 의사 레이블로 학습한 R1/R2의 전체 평가 PCK와 구분하여, 승인량, 정밀도 및 미검출을 포함한 전체 참조점 대비 유지율을 함께 해석하였다.')
    for p in io.role('ablation_figure'):
        if p.suffix=='.png':results.append({'type':'figure','path':str(p),'caption':io.files[str(p)].get('caption','통합 372장 필터 제거 비교')})
    results.append({'type':'h2','text':'4.4 재구성 참조에 대한 보조 자세 비교'})
    for kind in P.EXPECTED:
        rr=[r for r in pose if r['kind']==kind and r['condition'] in MAIN_CONDITIONS]
        table(results,kind+' 보조 자세 지표',['모델 R','PnP %','T 평균±SD cm','T n','번호 기준 R 평균±SD °','대칭 고려 R 평균±SD °','R n'],
            [[LABEL[r['condition']]+' '+r['round'],F.fmt(r['PnP_success_rate_full_DEV'],100),F.fmt(r['translation_error_cm_mean'])+' ± '+F.fmt(r['translation_error_cm_std']),r['translation_error_m_n'],
            F.fmt(r['rotation_indexed_deg_mean'])+' ± '+F.fmt(r['rotation_indexed_deg_std']),F.fmt(r['rotation_symmetry_deg_mean'])+' ± '+F.fmt(r['rotation_symmetry_deg_std']),r['rotation_indexed_deg_n']] for r in rr],
            [1.05,.55,1.15,.45,1.35,1.35,.45])
    paragraph(results,'보조 자세의 참조는 동일 인덱스의 2D 기준점, 등록 치수 및 카메라 행렬 K로부터 EPnP와 LM 보정을 사용하여 재구성하였다. 독립적으로 실측한 6D 정답은 아니다. PnP 성공률의 분모는 종류별 전체 평가 영상 수인 128/125/119장이며, 목재는 125장이다. 자세 오차는 참조와 모델의 추정이 모두 성공한 표본에 조건부로 계산하였다. 서로 다른 치수 가설 간 회전오차는 계산 불가로 처리하였다. 추가 80장의 카메라 보정 품질은 UNKNOWN이고, signed pose는 UNCONFIRMED 상태이다. 추가 617점은 수동 클릭 330점, PnP 투영 279점, 외삽 8점으로 구성된다. 목재의 고정 X 가설이 부적합했던 기록을 고려하여, 원래 규칙에 따라 재투영 RMSE가 최소인 보조 가설을 선택하는 정책을 고정하여 적용하였다.')
    results.append({'type':'h2','text':'4.5 실제·합성 자료 수 비율에 따른 민감도 분석'})
    table(results,'R2 고유 자료 수 비율의 15개 비교 조건',['종류 조건','실제 수','합성 수','합성/실제 %','PCK10 %','오차 평균±SD px','오차 n'],
        [[r['kind']+' '+r['condition'],r['real_unique'],r['synthetic_unique'],r['synthetic_percent'],F.fmt(r['PCK10'],100),F.fmt(r['conditional_pixel_error_mean'])+' ± '+F.fmt(r['conditional_pixel_error_std']),r['conditional_pixel_error_n']] for r in ratio],
        [1,.6,.6,.6,.8,1.75,.85])
    for kind in P.EXPECTED:
        rr=[r for r in ratio_ci if r['kind']==kind and r['metric']=='PCK10']
        table(results,kind+' 비율 모델의 PCK10 차이',['비교 모델','기준 모델','차이 %p','95% CI %p','판정'],
            [[r['candidate'],'합성 대조 R2' if r['comparator']=='control_R2' else r['comparator'],F.fmt(r['point_difference'],100),F.fmt(r['ci95_low'],100)+' ~ '+F.fmt(r['ci95_high'],100),decision_ko(r['decision'])] for r in rr],[.85,.95,.8,1.8,1.8])
    paragraph(results,'비율 실험을 실행하기 전에 목재 예비분석의 신뢰구간에 따른 통계적 중단 규칙을 적용하지 않기로 정하였다. 이에 따라 CI의 0 포함 여부와 관계없이 15개 조건을 모두 완료하였다. 이 사전 결정과 변경 이력은 부록의 재현 자료에 보존하였다. 체크포인트, 학습 시드 및 비교 조건은 고정하였으며, 성능 결과에 따라 모델이나 조건을 선택하지 않았다. 종류별 조건들은 동일한 고유 실제 자료와 고정된 R2 학습용 의사 레이블을 사용하므로, 조건별 차이를 학습용 의사 레이블 정밀도의 변화로 해석하지 않았다. 표의 합성/실제 비율은 실제 고유 자료 수 대비 합성 고유 자료 수의 백분율이다.')
    paragraph(results,'영상별 R2 학습 노출 횟수는 완료된 학습 기록의 ID를 집계하여 PER_FRAME_EXPOSURES.csv와 SUMMARY.csv에 제시하였다. 각 조건의 실제·합성 분기는 업데이트마다 각각 8회, 총 828개 업데이트에서 각각 6,624회 사용되었다. 고유 영상 수와 반복 노출 횟수를 구분하였으며, 영상당 평균과 모집단 SD, 최소·최대 노출 횟수는 부록에 정리하였다. 이 값은 R2의 추가학습 기록으로서 R1의 노출 횟수를 포함하지 않는다.')
    table(results,'학습을 마친 R2 모델의 가상 R3 교사 진단',['종류 조건 패널','승인 영상','선택 채점점','정밀도10 %','참조 유지율 %'],
        [[r['kind']+' '+r['condition']+' '+('전체 참조점' if r['panel']=='all_annotated' else '녹색 수동 600점'),r['accepted_frames'],r['scorable_accepted_points'],F.fmt(r['precision10'],100),F.fmt(r['selected_reference_coverage'],100)] for r in teacher],[2.1,.8,.9,1.2,1.2])
    paragraph(results,'이 표는 학습을 마친 R2를 다음 차수의 교사 모델로 사용한다고 가정했을 때의 필터 결과이며, 실제 R3 학습 결과는 아니다. 고정된 R1 오차 구간별 잔존오차는 기술 통계로 해석하였으며, 확증 편향의 원인을 입증하는 결과로 보지 않았다. 참조점의 혼합된 출처와 카메라 보정 품질 UNKNOWN에 따른 한계도 함께 고려하였다.')
    for p in io.role('ratio_figure'):
        if p.suffix=='.png':results.append({'type':'figure','path':str(p),'caption':io.files[str(p)].get('caption','고유 자료 수 비율 비교')})
    results.append({'type':'h2','text':'4.6 정성 비교 사례'})
    paragraph(results,'기존 DEV 정량 평가 이후 설명용으로 선정하여 공개한 6개 Input/R0/R1/R2 사례를 재사용하였다. 통합 목재 125장의 결과를 보고 사례를 다시 선택하지 않았다. 이 그림은 모델의 동작을 설명하는 예시이며, 모집단 평균이나 독립 실측 6D 정확도의 근거를 대체하지 않는다.')
    qualitative=io.role('qualitative_figure');P.require(len(qualitative)==6,'Exactly six frozen qualitative figures required')
    P.require(Counter(io.files[str(p)]['kind'] for p in qualitative)=={'PLASTIC':2,'WOOD':2,'GREEN':2},'Expected two fixed examples per type')
    for p in qualitative:
        item=io.files[str(p)];P.require(item.get('case_id'),'Explicit fixed case ID required')
        results.append({'type':'figure','path':str(p),'caption':item['kind']+' '+item['case_id']+' Input R0 R1 R2 비교'})
    discussion=chapter('제5장 고찰')
    for kind in P.EXPECTED:
        full,control=lookup[kind,'full',2],lookup[kind,'control',2]
        if full['available']=='True' and control['available']=='True':
            delta=P.number(full['PCK10'])-P.number(control['PCK10'])
            paragraph(discussion,f'{kind}의 전체 필터 R2에서 합성 대조 R2를 뺀 PCK10 차이는 {100*delta:.2f}%p였다. 기본 모델 대비 변화와 합성 추가학습 대비 변화를 구분하고, 짝지은 신뢰구간 및 관측률을 함께 고려해야 한다.')
    paragraph(discussion,'의사 레이블의 승인 수가 많다는 사실만으로 높은 정확도나 성능 개선이 보장되지는 않는다. 신뢰도가 낮은 점을 제외하면 남은 점의 조건부 정밀도는 높아질 수 있지만, 전체 참조점의 유지율과 관측률은 낮아질 수 있다. 따라서 승인 수, 정밀도 및 고정 참조점 유지율을 구분하여 제시하였다. 종류별로 배경, 재질 및 기준점의 출처가 다르므로 종류 간 성능 차이를 하나의 원인으로 설명하기는 어렵다.')
    paragraph(discussion,'본 실험은 단일 학습 시드와 반복 사용한 평가 자료에 기반한다. 혼합 해상도에서 원본 픽셀 문턱을 사용하였고, 목재의 일부 녹화는 학습 자료와 공유된다. 참조점은 수동·투영·외삽 출처가 혼합되어 있으며, 일부 카메라의 보정 품질은 UNKNOWN이다. 목재 녹화는 4개에 불과하고, 다른 종류에는 누락되거나 상충하는 녹화 메타데이터가 있다. 프레임 부트스트랩은 학습 시드에 따른 불확실성이나 독립 현장으로의 일반화 성능을 측정하지 않는다. 다중비교 보정은 수행하지 않았다.')
    paragraph(discussion,'R0와 R2의 차이에는 원학습과 현재 추가학습 사이의 입력 전처리 및 좌표 복호화 이력 차이가 포함된다. 따라서 이 변화를 의사 레이블만의 효과로 귀속하지 않았다. 현재 학습·평가 경로를 공유하는 합성 추가학습 대조군과 비율 조건 간 비교는 대상 효과를 구분하는 데 도움이 되지만, 독립 현장 일반화에 대한 인과적 검증으로 해석할 수는 없다.')
    conclusion=chapter('제6장 결론')
    paragraph(conclusion,'실제 파렛트의 의사 레이블을 이용하여 DOPE를 종류별로 자기학습하고, 기본 모델, 합성 추가학습 대조군, 필터 구성 요소 제거 조건 및 고유 자료 수 비율 조건을 동일한 372장에서 비교하였다. '+' '.join(result_phrase(k) for k in P.EXPECTED))
    paragraph(conclusion,'기본 모델 대비 변화와 함께 업데이트 수를 맞춘 대조군 및 짝지은 신뢰구간을 보고하였다. 재구성 참조에 대한 보조 자세 지표와 설명용 사례를 제공하였으나, 독립 실측 6D 정확도나 모든 조건의 우수성을 주장하지 않았다. 후속 검증에는 녹화와 장면을 분리한 자료, 카메라 정보 및 물리 모서리 규약을 확인한 독립 주석이 필요하다.')
    appendix=chapter('부록 재현 자료')
    paragraph(appendix,'전체 51개 평가조합의 2D 및 보조 자세 통계, 필터 제거 CI 180행, 기본 모델 간 비교 CI 90행, 비율 모델 15개와 CI 135행을 CSV로 제공한다. 수동 기준점 패널, 녹화 단위 CI, 가상 교사 모델 진단 및 영상별 노출 횟수도 함께 제공한다. 각 파일의 SHA와 실행·사전 결정·검증 이력을 보존하였다. 공개 묶음에는 설명용 입력 크롭과 비교 그림을 포함하되, 전체 원본 데이터셋, 주석 좌표, 체크포인트, 예측 배열 및 학습용 의사 레이블은 포함하지 않았다.')
    table(appendix,'비율 실험의 영상별 R2 학습 노출 횟수',['종류 조건','분기','고유 ID 수','총 노출 수','영상당 평균±SD','최소','최대'],
        [[r['kind']+' '+r['condition'],'실제' if r['branch']=='real' else '합성',r['unique_records'],r['total_draws'],
            F.fmt(r['mean_exposures'])+' ± '+F.fmt(r['population_SD_exposures']),r['min_exposures'],r['max_exposures']] for r in exposures],
        [1,.65,.65,.85,1.5,.6,.6])
    paragraph(appendix,'이 표의 SD는 영상별 노출 횟수의 모집단 표준편차이며, 학습 시드 간 변동이나 성능의 표준편차를 의미하지 않는다. 고유 수는 학습 기록의 ID 기준이며, 증강 후 픽셀의 동일성을 검사한 수가 아니다. 전체 영상별 횟수는 PER_FRAME_EXPOSURES.csv, 요약 통계는 SUMMARY.csv를 따른다.')
    return {'title':TITLE,'chapters':sections,'references':old['references'],'numerical_QA_sha256':P.sha(qa_path),'source_sha256':io.bindings}


def build(io,tag):
    import sys
    P.require(Path(sys.executable).resolve()==PYTHON.resolve(),'Use the authoritative bundled Python runtime')
    go_path,go,go_sha=manuscript_authority(io)
    data=content(io);io.unchanged();P.require(re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,63}',tag),'Unsafe manuscript version')
    folder=P.OUT/'manuscript_revision'/tag;P.require(not folder.exists(),'Preserve original/reviewed manuscript versions');folder.mkdir(parents=True)
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    marker('docx');doc=Document();sec=doc.sections[0]
    # Retain the original Korean thesis's A4 reference design, allowing natural flow.
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.left_margin=sec.right_margin=Inches(.85);sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.95)
    for name,size in [('Normal',11),('Title',20),('Heading 1',16),('Heading 2',13),('Caption',10)]:
        st=doc.styles[name];st.font.name='NanumMyeongjo';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'NanumMyeongjo')
        st.paragraph_format.line_spacing=1.4;st.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.widow_control=True
    doc.core_properties.title=TITLE;doc.core_properties.author='';doc.core_properties.subject='DOPE 종류별 자기학습 통합 372장 평가'
    doc.add_paragraph(TITLE,style='Title');doc.add_paragraph('통합 372장 평가와 필터 제거 및 고유 자료 수 비율 실험')
    heading_names=[]
    for c in data['chapters']:
        heading_names.append(c['title'])
        heading_names.extend('  '+b['text'] for b in c['blocks'] if b['type']=='h2' and b['text'].startswith('4.'))
    heading_names.append('참고문헌')
    doc.add_paragraph('목차',style='Heading 1')
    for title in heading_names:doc.add_paragraph(title)
    table_count=figure_count=0
    def mathrun(text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
    def sub(base,index):
        e=OxmlElement('m:sSub');b=OxmlElement('m:e');b.append(mathrun(base));s=OxmlElement('m:sub');s.append(mathrun(index));e.extend([b,s]);return e
    def equation(kind):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(30);p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
        om=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath')
        if kind=='loss':
            for e in (mathrun('L = '),sub('L','syn'),mathrun(' + λ'),sub('L','real'),mathrun(', λ = 1')):m.append(e)
        else:
            m.append(mathrun('PCK@τpx = '));f=OxmlElement('m:f');nu=OxmlElement('m:num');de=OxmlElement('m:den');nu.append(sub('N','correct'));nu.append(mathrun('(τ)'));de.append(sub('N','reference'));f.extend([nu,de]);m.append(f);m.append(mathrun(' × 100'))
        om.append(m);p._p.append(om)
    for chapter in data['chapters']:
        doc.add_page_break();doc.add_paragraph(chapter['title'],style='Heading 1')
        for b in chapter['blocks']:
            if b['type']=='p':doc.add_paragraph(b['text'])
            elif b['type']=='h2':doc.add_paragraph(b['text'],style='Heading 2')
            elif b['type']=='equation':equation(b['kind'])
            elif b['type']=='table':
                table_count+=1;caption=re.sub(r'^표\s*\d+\s*','',b['caption']);p=doc.add_paragraph(f'표 {table_count} {caption}',style='Caption');p.paragraph_format.keep_with_next=True
                t=doc.add_table(rows=1,cols=len(b['headers']));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
                widths=b.get('widths') or [6.55/len(b['headers'])]*len(b['headers']);scale=6.55/sum(widths);widths=[w*scale for w in widths]
                for col,width in zip(t.columns,widths):col.width=Inches(width)
                for c,w,label in zip(t.rows[0].cells,widths,b['headers']):c.width=Inches(w);c.text=str(label)
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for row in b['rows']:
                    for c,w,value in zip(t.add_row().cells,widths,row):c.width=Inches(w);c.text=str(value)
                borders=OxmlElement('w:tblBorders')
                for edge in ('top','left','bottom','right','insideH','insideV'):
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                t._tbl.tblPr.append(borders)
                for ri,row in enumerate(t.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell in row.cells:
                        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        margin=OxmlElement('w:tcMar')
                        for edge in ('top','bottom','left','right'):
                            e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margin.append(e)
                        cell._tc.get_or_add_tcPr().append(margin)
                        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E7EDF1' if ri==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shading)
                        for p in cell.paragraphs:
                            p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(0)
                            for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
                doc.add_paragraph()
            elif b['type']=='figure':
                figure_count+=1;path=io.path(b['path']);p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                from PIL import Image
                with Image.open(path) as image:iw,ih=image.size
                picture=p.add_run().add_picture(str(path),width=Inches(min(6.5,7.5*iw/ih)));picture._inline.docPr.set('descr',b['caption'])
                doc.add_paragraph(f'그림 {figure_count} '+b['caption'],style='Caption')
    doc.add_page_break();doc.add_paragraph('참고문헌',style='Heading 1')
    for ref in data['references']:doc.add_paragraph(ref)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;page=OxmlElement('w:fldSimple');page.set(qn('w:instr'),'PAGE');footer._p.append(page)
    for border in list(doc._element.xpath('.//w:pBdr'))+list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    target=folder/'DOPE_자기학습_통합372_졸업논문_초안.docx';doc.save(target)
    P.write(folder/'draft_content.json',data);io.unchanged();P.require(P.sha(go_path)==go_sha,'Manuscript GO changed');manuscript_authority(io)
    P.write(folder/'BUILD_RECEIPT.json',{'complete':True,'ready_for_delivery':False,'DOCX_sha256':P.sha(target),'source_sha256':io.bindings,
        'builder_sha256':go['builder_sha256'],'code_source_sha256':go['source_sha256'],'ROOT_GO_sha256':go_sha,'tables':table_count,'figures':figure_count,
        'page_budget':'natural flow; not forced18','toc':'chapter outline without guessed page numbers','all_pages_visual_QA':'PENDING'})
    return folder,target


def render(io,tag):
    P.require(re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_-]{0,63}',tag),'Unsafe manuscript tag')
    folder=P.OUT/'manuscript_revision'/tag;receipt=json.loads((folder/'BUILD_RECEIPT.json').read_text())
    target=folder/'DOPE_자기학습_통합372_졸업논문_초안.docx';P.require(receipt['DOCX_sha256']==P.sha(target),'Built DOCX changed')
    go_path,go,go_sha=manuscript_authority(io,receipt)
    wrapper=io.path(io.one('bundled_render_compat'))
    P.require(BUNDLED_LO.is_file(),'Bundled LibreOffice unavailable; never use system LO')
    marker('pdf');subprocess.run([str(PYTHON),'-B',str(wrapper),str(target),'--output_dir',str(folder/'render'),'--emit_pdf'],check=True)
    manuscript_authority(io,receipt);io.unchanged();P.require(go_sha==P.sha(go_path),'ROOT GO changed during render')
    P.write(folder/'RENDER_RECEIPT.json',{'render_command_completed':True,'all_pages_visual_QA':'PENDING','ready_for_delivery':False,
        'bundled_LibreOffice':str(BUNDLED_LO),'compatibility_wrapper_sha256':P.sha(wrapper),
        'rendered_files':P.inventory(folder/'render'),'DOCX_sha256':P.sha(target)})


def main():
    import sys
    P.require(Path(sys.executable).resolve()==PYTHON.resolve(),'Use bundled Python for document work')
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('stage',choices=('build','render'));parser.add_argument('--tag',required=True);args=parser.parse_args()
    io=P.Inputs(P.MANUSCRIPT_INPUTS);io.gates('numerical')
    if args.stage=='build':build(io,args.tag)
    else:render(io,args.tag)


if __name__=='__main__':main()
